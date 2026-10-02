from pathlib import Path
import hashlib

from sentence_transformers import SentenceTransformer
import chromadb
from pypdf import PdfReader
from docx import Document


# ==================================================
# 1. CONFIGURATION
# ==================================================

DATA_FOLDERS = [
    Path(r"C:\Users\srira\Documents"),
    Path(r"C:\Users\srira\Downloads"),
    Path(r"C:\Users\srira\Desktop"),
]


COLLECTION_NAME = "my_documents"


# IMPORTANT:
# True = rebuild the entire vector database
# False = incremental indexing
#
# Because we changed the chunking strategy,
# keep this TRUE for the first run.
#
REBUILD_COLLECTION = False


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".pdf",
    ".docx",
}


# ==================================================
# 2. EXCLUDED DIRECTORIES
# ==================================================

EXCLUDED_DIRECTORIES = {

    ".git",
    ".svn",
    ".hg",

    "venv",
    ".venv",
    "env",
    ".env",

    "node_modules",

    "__pycache__",

    ".pytest_cache",
    ".mypy_cache",

    "chroma_db",

    "$Recycle.Bin",
}


# ==================================================
# 3. EXCLUDED FILES
# ==================================================

EXCLUDED_FILES = {

    # Example:
    #
    # "passport.pdf",
    # "bank_statement.pdf",
    #

}


# ==================================================
# 4. CHECK IF FILE SHOULD BE EXCLUDED
# ==================================================

def should_exclude(file_path):

    # Check directory names
    for part in file_path.parts:

        if part in EXCLUDED_DIRECTORIES:

            return True


    # Check exact file name
    if file_path.name in EXCLUDED_FILES:

        return True


    return False


# ==================================================
# 5. READ TXT
# ==================================================

def read_txt(file_path):

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as file:

        return file.read()


# ==================================================
# 6. READ PDF
# ==================================================

def read_pdf(file_path):

    reader = PdfReader(
        str(file_path)
    )


    pages = []


    for page in reader.pages:

        text = page.extract_text()


        if text:

            pages.append(
                text
            )


    return "\n\n".join(
        pages
    )


# ==================================================
# 7. READ DOCX
# ==================================================

def read_docx(file_path):

    document = Document(
        str(file_path)
    )


    paragraphs = []


    for paragraph in document.paragraphs:

        text = paragraph.text.strip()


        if text:

            paragraphs.append(
                text
            )


    return "\n\n".join(
        paragraphs
    )


# ==================================================
# 8. READ FILE
# ==================================================

def read_file(file_path):

    extension = file_path.suffix.lower()


    if extension == ".txt":

        return read_txt(
            file_path
        )


    elif extension == ".pdf":

        return read_pdf(
            file_path
        )


    elif extension == ".docx":

        return read_docx(
            file_path
        )


    return None


# ==================================================
# 9. GENERIC TEXT CHUNKING
# ==================================================

def create_text_chunks(text):

    paragraphs = text.split(
        "\n\n"
    )


    chunks = []


    for paragraph in paragraphs:

        paragraph = paragraph.strip()


        if paragraph:

            chunks.append(
                paragraph
            )


    return chunks


# ==================================================
# 10. DOCX SMART CHUNKING
# ==================================================

def create_docx_chunks(file_path):

    document = Document(
        str(file_path)
    )


    chunks = []


    current_section = []


    for paragraph in document.paragraphs:

        text = paragraph.text.strip()


        if not text:

            continue


        # ------------------------------------------
        # Detect heading
        # ------------------------------------------

        style_name = ""

        if paragraph.style:

            style_name = (
                paragraph.style.name
                or ""
            )


        is_heading = (

            style_name.startswith(
                "Heading"
            )

            or text.endswith(":")

        )


        # ------------------------------------------
        # If this is a heading and we already
        # have a previous section, save it
        # ------------------------------------------

        if is_heading:

            if current_section:

                chunks.append(
                    "\n".join(
                        current_section
                    )
                )


                current_section = []


            current_section.append(
                text
            )


        else:

            current_section.append(
                text
            )


    # ------------------------------------------
    # Save final section
    # ------------------------------------------

    if current_section:

        chunks.append(
            "\n".join(
                current_section
            )
        )


    return chunks


# ==================================================
# 11. CREATE CHUNKS BASED ON FILE TYPE
# ==================================================

def create_chunks(
    file_path,
    content
):

    extension = file_path.suffix.lower()


    if extension == ".docx":

        return create_docx_chunks(
            file_path
        )


    return create_text_chunks(
        content
    )


# ==================================================
# 12. CALCULATE FILE HASH
# ==================================================

def calculate_file_hash(file_path):

    sha256 = hashlib.sha256()


    with open(
        file_path,
        "rb"
    ) as file:

        while chunk := file.read(
            4096
        ):

            sha256.update(
                chunk
            )


    return sha256.hexdigest()


# ==================================================
# 13. GET FULL SOURCE PATH
# ==================================================

def get_source_path(file_path):

    return str(
        file_path.resolve()
    ).replace(
        "\\",
        "/"
    )


# ==================================================
# 14. CREATE CHUNK ID
# ==================================================

def create_chunk_id(
    source,
    chunk_number
):

    source_hash = hashlib.sha256(
        source.encode("utf-8")
    ).hexdigest()


    return (
        f"{source_hash}_{chunk_number}"
    )


# ==================================================
# 15. LOAD EMBEDDING MODEL
# ==================================================

print(
    "\nLoading embedding model..."
)


model = SentenceTransformer(
    "all-MiniLM-L6-v2"
)


# ==================================================
# 16. CONNECT TO CHROMADB
# ==================================================

client = chromadb.PersistentClient(
    path="./chroma_db"
)


# ==================================================
# 17. CREATE / LOAD COLLECTION
# ==================================================

if REBUILD_COLLECTION:

    print(
        "\nRebuilding collection..."
    )


    try:

        client.delete_collection(
            COLLECTION_NAME
        )

        print(
            "Old collection deleted."
        )

    except Exception:

        print(
            "No existing collection found."
        )


    collection = client.create_collection(
        COLLECTION_NAME
    )


    print(
        "New collection created."
    )


else:

    try:

        collection = client.get_collection(
            COLLECTION_NAME
        )


        print(
            "Existing collection loaded."
        )


    except Exception:

        collection = client.create_collection(
            COLLECTION_NAME
        )


        print(
            "New collection created."
        )


# ==================================================
# 18. READ EXISTING VECTOR DATA
# ==================================================

existing_data = collection.get(
    include=[
        "metadatas"
    ]
)


existing_files = {}


for metadata in existing_data["metadatas"]:

    source = metadata.get(
        "source"
    )


    file_hash = metadata.get(
        "file_hash"
    )


    if source and file_hash:

        existing_files[
            source
        ] = file_hash


# ==================================================
# 19. TRACK CURRENT FILES
# ==================================================

current_files = set()


# ==================================================
# 20. SCAN ALL DATA FOLDERS
# ==================================================

for data_folder in DATA_FOLDERS:

    print(
        f"\nScanning folder: {data_folder}"
    )


    if not data_folder.exists():

        print(
            f"WARNING: Folder does not exist: "
            f"{data_folder}"
        )

        continue


    for file_path in data_folder.rglob("*"):


        # ------------------------------------------
        # Only files
        # ------------------------------------------

        if not file_path.is_file():

            continue


        # ------------------------------------------
        # Exclusions
        # ------------------------------------------

        if should_exclude(
            file_path
        ):

            continue


        # ------------------------------------------
        # Supported file types
        # ------------------------------------------

        extension = (
            file_path.suffix.lower()
        )


        if extension not in SUPPORTED_EXTENSIONS:

            continue


        # ------------------------------------------
        # Full source path
        # ------------------------------------------

        source = get_source_path(
            file_path
        )


        current_files.add(
            source
        )


        print(
            f"\nChecking: {source}"
        )


        # ------------------------------------------
        # Calculate hash
        # ------------------------------------------

        try:

            current_hash = (
                calculate_file_hash(
                    file_path
                )
            )

        except Exception as error:

            print(
                f"ERROR hashing file: {error}"
            )

            continue


        # ==================================================
        # NEW FILE
        # ==================================================

        if source not in existing_files:

            print(
                "Status: NEW"
            )


            try:

                content = read_file(
                    file_path
                )

            except Exception as error:

                print(
                    f"ERROR reading file: {error}"
                )

                continue


            if not content:

                print(
                    "WARNING: No text extracted."
                )

                continue


            chunks = create_chunks(
                file_path,
                content
            )


            print(
                f"Chunks created: {len(chunks)}"
            )


            for number, chunk in enumerate(
                chunks,
                start=1
            ):


                embedding = model.encode(
                    chunk
                ).tolist()


                chunk_id = create_chunk_id(
                    source,
                    number
                )


                collection.add(

                    ids=[
                        chunk_id
                    ],

                    documents=[
                        chunk
                    ],

                    embeddings=[
                        embedding
                    ],

                    metadatas=[

                        {
                            "source": source,
                            "chunk": number,
                            "file_hash": current_hash,
                            "file_type": extension
                        }

                    ]
                )


            continue


        # ==================================================
        # EXISTING FILE
        # ==================================================

        old_hash = existing_files[
            source
        ]


        # ==================================================
        # UNCHANGED
        # ==================================================

        if current_hash == old_hash:

            print(
                "Status: UNCHANGED"
            )

            continue


        # ==================================================
        # MODIFIED
        # ==================================================

        print(
            "Status: MODIFIED"
        )


        old_chunks = collection.get(

            where={
                "source": source
            },

            include=[
                "metadatas"
            ]
        )


        old_ids = old_chunks[
            "ids"
        ]


        if old_ids:

            collection.delete(
                ids=old_ids
            )


        try:

            content = read_file(
                file_path
            )

        except Exception as error:

            print(
                f"ERROR reading file: {error}"
            )

            continue


        if not content:

            print(
                "WARNING: No text extracted."
            )

            continue


        chunks = create_chunks(
            file_path,
            content
        )


        print(
            f"Chunks created: {len(chunks)}"
        )


        for number, chunk in enumerate(
            chunks,
            start=1
        ):


            embedding = model.encode(
                chunk
            ).tolist()


            chunk_id = create_chunk_id(
                source,
                number
            )


            collection.add(

                ids=[
                    chunk_id
                ],

                documents=[
                    chunk
                ],

                embeddings=[
                    embedding
                ],

                metadatas=[

                    {
                        "source": source,
                        "chunk": number,
                        "file_hash": current_hash,
                        "file_type": extension
                    }

                ]
            )


# ==================================================
# 21. DELETE VECTORS FOR DELETED FILES
# ==================================================

for source in existing_files:

    if source not in current_files:

        print(
            f"\nStatus: DELETED - {source}"
        )


        deleted_chunks = collection.get(

            where={
                "source": source
            },

            include=[
                "metadatas"
            ]
        )


        deleted_ids = deleted_chunks[
            "ids"
        ]


        if deleted_ids:

            collection.delete(
                ids=deleted_ids
            )


# ==================================================
# 22. FINAL RESULT
# ==================================================

print(
    "\n========================================"
)

print(
    "Incremental ingestion complete."
)

print(
    "Total chunks:",
    collection.count()
)

print(
    "========================================"
)