from pathlib import Path
import hashlib
import uuid

import chromadb
from sentence_transformers import SentenceTransformer

# OCR / PDF
from pypdf import PdfReader
from pdf2image import convert_from_path
import pytesseract

# DOCX
from docx import Document


# ============================================================
# CONFIGURATION
# ============================================================

COLLECTION_NAME = "my_documents"
CHROMA_PATH = "./chroma_db"

OCR_ENABLED = True
OCR_DPI = 200

SUPPORTED_EXTENSIONS = {
    ".txt",
    ".pdf",
    ".docx",
}


# ============================================================
# LOAD MODELS / DATABASE
# ============================================================

print("Loading embedding model...")

embedding_model = SentenceTransformer(
    "all-MiniLM-L6-v2"
)

print("Embedding model loaded.")

chroma_client = chromadb.PersistentClient(
    path=CHROMA_PATH
)

collection = chroma_client.get_or_create_collection(
    name=COLLECTION_NAME
)


# ============================================================
# PATH NORMALIZATION
# ============================================================

def normalize_path(path):
    """
    Normalize Windows paths so the same file always
    gets the same source value.
    """

    return str(
        Path(path).resolve()
    ).replace("\\", "/").lower()


# ============================================================
# FILE HASH
# ============================================================

def calculate_file_hash(file_path):
    """
    Calculate SHA256 hash of a file.

    Used to identify the exact version of a document.
    """

    sha256 = hashlib.sha256()

    with open(file_path, "rb") as f:

        while True:

            data = f.read(1024 * 1024)

            if not data:
                break

            sha256.update(data)

    return sha256.hexdigest()


# ============================================================
# TXT READER
# ============================================================

def read_txt(file_path):

    encodings = [
        "utf-8",
        "utf-8-sig",
        "cp1252",
        "latin-1",
    ]

    for encoding in encodings:

        try:

            with open(
                file_path,
                "r",
                encoding=encoding,
                errors="ignore",
            ) as f:

                return f.read()

        except Exception:
            continue

    return ""


# ============================================================
# PDF READER
# ============================================================

def read_pdf(file_path):

    text_parts = []

    try:

        reader = PdfReader(file_path)

        for page in reader.pages:

            try:

                text = page.extract_text()

                if text:
                    text_parts.append(text)

            except Exception as e:

                print(
                    f"PDF page extraction warning: {e}"
                )

    except Exception as e:

        print(
            f"PDF extraction failed: {e}"
        )

    normal_text = "\n".join(text_parts).strip()

    # --------------------------------------------------------
    # If normal PDF extraction worked, use it.
    # --------------------------------------------------------

    if normal_text:

        return normal_text

    # --------------------------------------------------------
    # OCR fallback for scanned PDFs
    # --------------------------------------------------------

    if not OCR_ENABLED:

        return ""

    print(
        f"OCR required: {Path(file_path).name}"
    )

    try:

        pages = convert_from_path(
            file_path,
            dpi=OCR_DPI,
        )

        ocr_text = []

        for index, page in enumerate(pages, start=1):

            print(
                f"  OCR page {index}/{len(pages)}"
            )

            text = pytesseract.image_to_string(
                page
            )

            if text:

                ocr_text.append(text)

        return "\n".join(ocr_text).strip()

    except Exception as e:

        print(
            f"OCR failed: {e}"
        )

        return ""


# ============================================================
# DOCX READER
# ============================================================

def read_docx(file_path):

    try:

        document = Document(file_path)

        sections = []

        current_heading = None
        current_content = []

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if not text:
                continue

            style_name = paragraph.style.name.lower()

            is_heading = (
                "heading" in style_name
            )

            if is_heading:

                # Save previous section
                if current_heading and current_content:

                    sections.append(
                        current_heading
                        + "\n"
                        + "\n".join(current_content)
                    )

                current_heading = text
                current_content = []

            else:

                current_content.append(text)

        # Save final section
        if current_heading and current_content:

            sections.append(
                current_heading
                + "\n"
                + "\n".join(current_content)
            )

        # If no headings existed
        if not sections:

            all_text = []

            for paragraph in document.paragraphs:

                text = paragraph.text.strip()

                if text:

                    all_text.append(text)

            return "\n".join(all_text)

        return "\n\n".join(sections)

    except Exception as e:

        print(
            f"DOCX extraction failed: {e}"
        )

        return ""


# ============================================================
# FILE READER
# ============================================================

def read_file(file_path):

    file_path = Path(file_path)

    extension = file_path.suffix.lower()

    try:

        if extension == ".txt":

            return read_txt(file_path)

        elif extension == ".pdf":

            return read_pdf(file_path)

        elif extension == ".docx":

            return read_docx(file_path)

        else:

            print(
                f"Unsupported file type: {extension}"
            )

            return ""

    except Exception as e:

        print(
            f"Failed to read {file_path}: {e}"
        )

        return ""


# ============================================================
# CHUNKING
# ============================================================

def create_text_chunks(
    text,
    chunk_size=1000,
    overlap=150,
):
    """
    Simple overlapping chunking.

    Example:

    chunk 1: characters 0-1000
    chunk 2: characters 850-1850
    chunk 3: characters 1700-2700
    """

    text = text.strip()

    if not text:

        return []

    chunks = []

    start = 0

    text_length = len(text)

    while start < text_length:

        end = start + chunk_size

        chunk = text[start:end].strip()

        if chunk:

            chunks.append(chunk)

        if end >= text_length:

            break

        start = end - overlap

    return chunks


# ============================================================
# CREATE RECORDS
# ============================================================

def create_records(
    file_path,
    chunks,
    file_hash,
):

    normalized_path = normalize_path(
        file_path
    )

    ids = []
    documents = []
    metadatas = []

    for index, chunk in enumerate(chunks):

        chunk_id = str(
            uuid.uuid4()
        )

        ids.append(chunk_id)

        documents.append(chunk)

        metadatas.append(
            {
                "source": normalized_path,
                "filename": Path(file_path).name,
                "extension": Path(file_path).suffix.lower(),
                "chunk_index": index,
                "file_hash": file_hash,
            }
        )

    return ids, documents, metadatas


# ============================================================
# DELETE FILE FROM CHROMA
# ============================================================

def delete_file(file_path):

    normalized_path = normalize_path(
        file_path
    )

    print(
        f"Removing from Chroma: {normalized_path}"
    )

    try:

        existing = collection.get(
            where={
                "source": normalized_path
            }
        )

        existing_ids = existing.get(
            "ids",
            []
        )

        if existing_ids:

            collection.delete(
                ids=existing_ids
            )

            print(
                f"Removed {len(existing_ids)} chunks."
            )

        else:

            print(
                "No Chroma records found."
            )

    except Exception as e:

        print(
            f"Delete failed: {e}"
        )


# ============================================================
# INGEST FILE
# ============================================================

def ingest_file(file_path):

    file_path = Path(file_path)

    print(
        f"\nINGESTING: {file_path.name}"
    )

    # --------------------------------------------------------
    # Check existence
    # --------------------------------------------------------

    if not file_path.exists():

        print(
            "File does not exist."
        )

        return False

    # --------------------------------------------------------
    # Check extension
    # --------------------------------------------------------

    extension = file_path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:

        print(
            f"Unsupported extension: {extension}"
        )

        return False

    # --------------------------------------------------------
    # Calculate hash
    # --------------------------------------------------------

    try:

        file_hash = calculate_file_hash(
            file_path
        )

    except Exception as e:

        print(
            f"Hash calculation failed: {e}"
        )

        return False

    # --------------------------------------------------------
    # Read document
    # --------------------------------------------------------

    try:

        text = read_file(
            file_path
        )

    except Exception as e:

        print(
            f"Reading failed: {e}"
        )

        return False

    if not text.strip():

        print(
            "No text found. File was not indexed."
        )

        return False

    # --------------------------------------------------------
    # Create chunks
    # --------------------------------------------------------

    chunks = create_text_chunks(
        text
    )

    if not chunks:

        print(
            "No chunks created."
        )

        return False

    print(
        f"Created {len(chunks)} chunks."
    )

    # --------------------------------------------------------
    # Create Chroma records
    # --------------------------------------------------------

    ids, documents, metadatas = create_records(
        file_path,
        chunks,
        file_hash,
    )

    # --------------------------------------------------------
    # Generate embeddings FIRST
    # --------------------------------------------------------

    print(
        "Generating embeddings..."
    )

    try:

        embeddings = embedding_model.encode(
            documents,
            show_progress_bar=False,
        )

    except Exception as e:

        print(
            f"Embedding generation failed: {e}"
        )

        print(
            "OLD VERSION WAS PRESERVED."
        )

        return False

    # --------------------------------------------------------
    # Store NEW version FIRST
    # --------------------------------------------------------

    print(
        "Storing new version..."
    )

    try:

        collection.add(
            ids=ids,
            documents=documents,
            metadatas=metadatas,
            embeddings=embeddings.tolist(),
        )

    except Exception as e:

        print(
            f"Chroma storage failed: {e}"
        )

        print(
            "OLD VERSION WAS PRESERVED."
        )

        return False

    # --------------------------------------------------------
    # NEW VERSION SUCCESSFULLY STORED
    # --------------------------------------------------------

    print(
        "New version stored successfully."
    )

    # --------------------------------------------------------
    # Remove OLD chunks
    #
    # IMPORTANT:
    # We now remove old chunks only AFTER
    # successful storage.
    # --------------------------------------------------------

    normalized_path = normalize_path(
        file_path
    )

    try:

        existing = collection.get(
            where={
                "source": normalized_path
            }
        )

        existing_ids = existing.get(
            "ids",
            []
        )

        # The query now also contains the newly
        # inserted records, so keep the new IDs.

        old_ids = [
            item_id
            for item_id in existing_ids
            if item_id not in ids
        ]

        if old_ids:

            print(
                f"Removing {len(old_ids)} old chunks..."
            )

            collection.delete(
                ids=old_ids
            )

    except Exception as e:

        print(
            f"Warning: old-version cleanup failed: {e}"
        )

        print(
            "The new version is still available."
        )

    print(
        f"Successfully indexed {len(chunks)} chunks."
    )

    return True


# ============================================================
# MODULE TEST
# ============================================================

if __name__ == "__main__":

    print()
    print(
        "Ingestion module loaded successfully."
    )

    print(
        "Supported extensions:",
        SUPPORTED_EXTENSIONS,
    )

    print(
        "OCR enabled:",
        OCR_ENABLED,
    )

    print(
        "Chroma collection:",
        COLLECTION_NAME,
    )

    print(
        "Ready for the folder watcher."
    )